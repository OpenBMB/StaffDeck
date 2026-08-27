from __future__ import annotations

import hashlib
import importlib
from io import BytesIO

import pytest
from docx import Document


def _load_extraction_module():
    return importlib.import_module("app.documents.extraction")


class _FakePdfPage:
    def __init__(self, text: str) -> None:
        self._text = text

    def extract_text(self) -> str:
        return self._text


def test_native_pdf_returns_full_contract_when_every_page_has_text(monkeypatch: pytest.MonkeyPatch) -> None:
    extraction = _load_extraction_module()

    class FakeReader:
        def __init__(self) -> None:
            self.pages = [_FakePdfPage("First page"), _FakePdfPage("Second page")]

    monkeypatch.setattr("pypdf.PdfReader", lambda _stream: FakeReader())
    content = b"%PDF-1.4 native"

    result = extraction.DocumentExtractor().extract("policy.pdf", content)

    assert result.method == "native"
    assert result.engine == "pypdf"
    assert result.source_sha256 == hashlib.sha256(content).hexdigest()
    assert result.source_page_count == 2
    assert result.page_refs == ["page:1", "page:2"]
    assert result.char_count == len("First pageSecond page")
    assert result.table_count == 0
    assert result.pages[0].page_number == 1
    assert result.pages[0].char_count == len("First page")
    assert result.pages[1].text == "Second page"
    assert "# PDF 文档" in result.text
    assert "## 第 2 页" in result.text


def test_scanned_pdf_uses_injected_structured_adapter_when_enabled(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    extraction = _load_extraction_module()

    class FakeReader:
        def __init__(self) -> None:
            self.pages = [_FakePdfPage("Native page"), _FakePdfPage("")]

    class FakeAdapter:
        def __init__(self) -> None:
            self.calls: list[tuple[str, bytes]] = []

        def extract_pdf(self, filename: str, content: bytes):
            self.calls.append((filename, content))
            return extraction.DocumentExtractionResult(
                text="structured text",
                pages=[
                    extraction.ExtractedPage(
                        page_number=1,
                        text="structured text",
                        char_count=len("structured text"),
                        table_count=1,
                    )
                ],
                page_refs=["page:1"],
                source_sha256=hashlib.sha256(content).hexdigest(),
                source_page_count=1,
                method="structured",
                engine="rapiddoc-ort",
                engine_version="0.0-test",
                warnings=["OCR_USED"],
                char_count=len("structured text"),
                table_count=1,
            )

    monkeypatch.setattr("pypdf.PdfReader", lambda _stream: FakeReader())
    adapter = FakeAdapter()
    content = b"%PDF-1.4 scanned"

    result = extraction.DocumentExtractor(
        structured_pdf_enabled=True,
        structured_pdf_adapter=adapter,
    ).extract("scan.pdf", content)

    assert adapter.calls == [("scan.pdf", content)]
    assert result.method == "structured"
    assert result.engine == "rapiddoc-ort"
    assert result.warnings == ["OCR_USED"]


def test_scanned_pdf_without_adapter_raises_stable_error_code(monkeypatch: pytest.MonkeyPatch) -> None:
    extraction = _load_extraction_module()

    class FakeReader:
        def __init__(self) -> None:
            self.pages = [_FakePdfPage(""), _FakePdfPage("also empty")]

    monkeypatch.setattr("pypdf.PdfReader", lambda _stream: FakeReader())

    with pytest.raises(extraction.DocumentExtractionError) as exc_info:
        extraction.DocumentExtractor(structured_pdf_enabled=True).extract("scan.pdf", b"%PDF-1.4")

    assert exc_info.value.code == "OCR_DEPENDENCY_MISSING"


@pytest.mark.parametrize(
    "code",
    [
        "OCR_MODEL_MISSING",
        "DOCUMENT_EXTRACTION_FAILED",
        "OCR_TIMEOUT",
        "PDF_CORRUPTED",
        "PDF_ENCRYPTED",
    ],
)
def test_structured_adapter_errors_surface_stable_codes(
    monkeypatch: pytest.MonkeyPatch,
    code: str,
) -> None:
    extraction = _load_extraction_module()

    class FakeReader:
        def __init__(self) -> None:
            self.pages = [_FakePdfPage(""), _FakePdfPage("still empty")]

    class FakeAdapter:
        def extract_pdf(self, filename: str, content: bytes):
            raise extraction.DocumentExtractionError(code, f"{code}: boom")

    monkeypatch.setattr("pypdf.PdfReader", lambda _stream: FakeReader())

    with pytest.raises(extraction.DocumentExtractionError) as exc_info:
        extraction.DocumentExtractor(
            structured_pdf_enabled=True,
            structured_pdf_adapter=FakeAdapter(),
        ).extract("scan.pdf", b"%PDF-1.4")

    assert exc_info.value.code == code
    assert str(exc_info.value).startswith(code)


@pytest.mark.parametrize(
    ("filename", "payload", "expected_text"),
    [
        ("note.txt", b"plain text payload", "plain text payload"),
        ("note.md", b"# title\n\nbody", "# title\n\nbody"),
    ],
)
def test_text_and_markdown_continue_to_use_native_parsers(
    filename: str,
    payload: bytes,
    expected_text: str,
) -> None:
    extraction = _load_extraction_module()

    result = extraction.DocumentExtractor().extract(filename, payload)

    assert result.method == "native"
    assert result.source_page_count == 1
    assert result.pages[0].text == expected_text
    assert result.text == expected_text


def test_docx_continues_to_use_existing_parser() -> None:
    extraction = _load_extraction_module()
    document = Document()
    document.add_heading("Policy", level=1)
    document.add_paragraph("First paragraph")
    payload = BytesIO()
    document.save(payload)

    result = extraction.DocumentExtractor().extract("policy.docx", payload.getvalue())

    assert result.method == "native"
    assert result.engine == "python-docx"
    assert result.pages[0].text.startswith("# Policy")
    assert "First paragraph" in result.text


def test_core_result_is_not_truncated_and_only_preview_helper_truncates() -> None:
    extraction = _load_extraction_module()
    long_text = "A" * 25010

    result = extraction.DocumentExtractor().extract("long.txt", long_text.encode("utf-8"))

    assert result.text == long_text
    assert result.pages[0].text == long_text
    assert result.char_count == len(long_text)
    assert result.preview() == ("A" * 24000) + "..."


def test_legacy_extract_text_keeps_tuple_contract_for_pdf(monkeypatch: pytest.MonkeyPatch) -> None:
    from app.knowledge.parser import extract_text

    class FakeReader:
        def __init__(self) -> None:
            self.pages = [_FakePdfPage("Alpha"), _FakePdfPage("Beta")]

    monkeypatch.setattr("pypdf.PdfReader", lambda _stream: FakeReader())

    text, file_type = extract_text("policy.pdf", b"%PDF-1.4")

    assert file_type == "pdf"
    assert text.startswith("# PDF 文档")
    assert "Alpha" in text
    assert "Beta" in text


def test_legacy_extract_text_uses_configured_structured_extractor_for_scanned_pdf(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.config import Settings
    from app.knowledge import parser

    class FakeReader:
        def __init__(self) -> None:
            self.pages = [_FakePdfPage("native text"), _FakePdfPage("")]

    class FakeAdapter:
        def __init__(
            self,
            *,
            model_manager,
            timeout_seconds: float,
            worker_count: int,
            max_pages: int,
            max_pixels: int,
            device: str = "cpu",
            engine: str = "ort",
        ) -> None:
            self.model_manager = model_manager
            self.timeout_seconds = timeout_seconds
            self.worker_count = worker_count
            self.max_pages = max_pages
            self.max_pixels = max_pixels
            self.device = device
            self.engine = engine
            self.calls: list[tuple[str, bytes]] = []

        def extract_pdf(self, filename: str, content: bytes):
            self.calls.append((filename, content))
            return extraction.DocumentExtractionResult(
                text="structured route",
                pages=[
                    extraction.ExtractedPage(
                        page_number=1,
                        text="structured route",
                        char_count=len("structured route"),
                    )
                ],
                page_refs=["page:1"],
                source_sha256=hashlib.sha256(content).hexdigest(),
                source_page_count=1,
                method="structured",
                engine="rapiddoc-ort",
                engine_version="test",
                warnings=[],
                char_count=len("structured route"),
                table_count=0,
            )

    extraction = _load_extraction_module()
    fake_adapter_instances: list[FakeAdapter] = []

    def build_fake_adapter(**kwargs):
        adapter = FakeAdapter(**kwargs)
        fake_adapter_instances.append(adapter)
        return adapter

    settings = Settings().model_copy(
        update={
            "structured_pdf_enabled": True,
            "structured_pdf_engine": "rapiddoc",
            "rapid_models_dir": "C:/configured-models",
            "structured_pdf_timeout_seconds": 12.5,
            "structured_pdf_worker_count": 3,
            "structured_pdf_max_pages": 9,
            "structured_pdf_max_pixels": 123456,
        }
    )

    monkeypatch.setattr("pypdf.PdfReader", lambda _stream: FakeReader())
    monkeypatch.setattr(parser, "get_settings", lambda: settings)
    monkeypatch.setattr(parser, "RapidDocStructuredPdfAdapter", build_fake_adapter)

    text, file_type = parser.extract_text("scan.pdf", b"%PDF-1.4 scanned")

    assert file_type == "pdf"
    assert text == "structured route"
    assert len(fake_adapter_instances) == 1
    adapter = fake_adapter_instances[0]
    assert adapter.calls == [("scan.pdf", b"%PDF-1.4 scanned")]
    assert str(adapter.model_manager.resolve_model_dir()) == "C:\\configured-models"
    assert adapter.timeout_seconds == 12.5
    assert adapter.worker_count == 3
    assert adapter.max_pages == 9
    assert adapter.max_pixels == 123456
