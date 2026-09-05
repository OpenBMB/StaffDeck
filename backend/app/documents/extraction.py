from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO
from pathlib import Path
from typing import Protocol
from zipfile import ZipFile

SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".html", ".htm", ".pdf", ".docx", ".doc"}


class StructuredPdfAdapter(Protocol):
    def extract_pdf(self, filename: str, content: bytes) -> DocumentExtractionResult: ...


class DocumentExtractionError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        self.code = code
        self.message = message
        super().__init__(f"{code}: {message}")


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int
    text: str
    char_count: int
    table_count: int = 0

    @property
    def ref(self) -> str:
        return f"page:{self.page_number}"


@dataclass(frozen=True)
class DocumentExtractionResult:
    text: str
    pages: list[ExtractedPage]
    page_refs: list[str]
    source_sha256: str
    source_page_count: int
    method: str
    engine: str
    engine_version: str
    warnings: list[str] = field(default_factory=list)
    char_count: int = 0
    table_count: int = 0

    def preview(self, limit: int = 24000) -> str:
        if limit < 0:
            raise ValueError("limit must be non-negative")
        if len(self.text) <= limit:
            return self.text
        return self.text[:limit] + "..."


class DocumentExtractor:
    def __init__(
        self,
        *,
        structured_pdf_enabled: bool = False,
        structured_pdf_adapter: StructuredPdfAdapter | None = None,
    ) -> None:
        self.structured_pdf_enabled = structured_pdf_enabled
        self.structured_pdf_adapter = structured_pdf_adapter

    def extract(self, filename: str, content: bytes) -> DocumentExtractionResult:
        suffix = Path(filename).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_FAILED",
                f"unsupported file format: {suffix or 'unknown'}",
            )
        if suffix == ".doc":
            raise DocumentExtractionError(
                "DOCUMENT_EXTRACTION_FAILED",
                "legacy .doc is unsupported; convert to .docx first",
            )
        if suffix in {".txt", ".md", ".markdown"}:
            text = _decode_text(content)
            return _single_page_result(
                text=text,
                source_bytes=content,
                source_page_count=1,
                method="native",
                engine="text",
                engine_version="builtin",
            )
        if suffix in {".html", ".htm"}:
            return _single_page_result(
                text=_extract_html(content),
                source_bytes=content,
                source_page_count=1,
                method="native",
                engine="html-parser",
                engine_version="builtin",
            )
        if suffix == ".docx":
            text, engine = _extract_docx(content)
            return _single_page_result(
                text=text,
                source_bytes=content,
                source_page_count=1,
                method="native",
                engine=engine,
                engine_version="builtin",
            )
        if suffix == ".pdf":
            return self._extract_pdf(filename, content)
        raise DocumentExtractionError("DOCUMENT_EXTRACTION_FAILED", f"unsupported file format: {suffix}")

    def _extract_pdf(self, filename: str, content: bytes) -> DocumentExtractionResult:
        try:
            from pypdf import PdfReader
            from pypdf.errors import PdfReadError
        except Exception as exc:  # pragma: no cover - environment dependent
            raise DocumentExtractionError("DOCUMENT_EXTRACTION_FAILED", "missing pypdf") from exc

        try:
            reader = PdfReader(BytesIO(content))
        except PdfReadError as exc:
            raise DocumentExtractionError("PDF_CORRUPTED", str(exc) or "pdf is corrupted") from exc
        except Exception as exc:
            raise DocumentExtractionError("DOCUMENT_EXTRACTION_FAILED", str(exc) or "failed to read pdf") from exc

        if getattr(reader, "is_encrypted", False):
            raise DocumentExtractionError("PDF_ENCRYPTED", "pdf is encrypted")

        page_texts: list[str] = []
        for page in reader.pages:
            try:
                page_texts.append(page.extract_text() or "")
            except Exception as exc:
                raise DocumentExtractionError("DOCUMENT_EXTRACTION_FAILED", str(exc) or "failed to extract pdf text") from exc

        if page_texts and all(text.strip() for text in page_texts):
            pages = [
                ExtractedPage(page_number=index + 1, text=text, char_count=len(text))
                for index, text in enumerate(page_texts)
            ]
            return DocumentExtractionResult(
                text=_render_pdf_text(pages),
                pages=pages,
                page_refs=[page.ref for page in pages],
                source_sha256=_sha256_bytes(content),
                source_page_count=len(page_texts),
                method="native",
                engine="pypdf",
                engine_version=_package_version("pypdf"),
                warnings=[],
                char_count=sum(page.char_count for page in pages),
                table_count=0,
            )

        if not self.structured_pdf_enabled or self.structured_pdf_adapter is None:
            raise DocumentExtractionError(
                "OCR_DEPENDENCY_MISSING",
                "structured pdf adapter is not configured",
            )

        return self.structured_pdf_adapter.extract_pdf(filename, content)


def _single_page_result(
    *,
    text: str,
    source_bytes: bytes,
    source_page_count: int,
    method: str,
    engine: str,
    engine_version: str,
) -> DocumentExtractionResult:
    page = ExtractedPage(page_number=1, text=text, char_count=len(text))
    return DocumentExtractionResult(
        text=text,
        pages=[page],
        page_refs=[page.ref],
        source_sha256=_sha256_bytes(source_bytes),
        source_page_count=source_page_count,
        method=method,
        engine=engine,
        engine_version=engine_version,
        warnings=[],
        char_count=len(text),
        table_count=0,
    )


def _render_pdf_text(pages: list[ExtractedPage]) -> str:
    if not pages:
        return ""
    body = "\n\n".join(f"## 第 {page.page_number} 页\n\n{page.text}" for page in pages if page.text.strip())
    if not body:
        return ""
    return "# PDF 文档\n\n" + body


def _sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _package_version(package_name: str) -> str:
    try:
        return version(package_name)
    except PackageNotFoundError:
        return "unknown"


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


def _extract_html(content: bytes) -> str:
    text = _decode_text(content)
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(text, "html.parser")
        for item in soup(["script", "style", "noscript"]):
            item.decompose()
        return soup.get_text("\n")
    except (ImportError, TypeError, ValueError):
        parser = _HTMLTextExtractor()
        parser.feed(text)
        return parser.text


def _extract_docx(content: bytes) -> tuple[str, str]:
    try:
        from docx import Document

        document = Document(BytesIO(content))
        rows: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            heading_level = _docx_heading_level(paragraph.style.name if paragraph.style else "")
            rows.append(f"{'#' * heading_level} {text}" if heading_level else text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    rows.append(" | ".join(cells))
        return "\n".join(rows), "python-docx"
    except (ImportError, KeyError, TypeError, ValueError):
        return _extract_docx_with_zip(content), "docx-zip"


def _docx_heading_level(style_name: str) -> int | None:
    match = re.match(r"^(?:Heading|标题)\s*([1-6])$", style_name.strip(), re.IGNORECASE)
    return int(match.group(1)) if match else None


def _extract_docx_with_zip(content: bytes) -> str:
    try:
        with ZipFile(BytesIO(content)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception as exc:
        raise DocumentExtractionError("DOCUMENT_EXTRACTION_FAILED", "unable to parse docx") from exc
    parser = _DocxTextExtractor()
    parser.feed(xml)
    return parser.text


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []

    @property
    def text(self) -> str:
        return "\n".join(part.strip() for part in self._parts if part.strip())

    def handle_data(self, data: str) -> None:
        if data.strip():
            self._parts.append(data)


class _DocxTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._parts: list[str] = []

    @property
    def text(self) -> str:
        return "\n".join(part.strip() for part in self._parts if part.strip())

    def handle_data(self, data: str) -> None:
        if data.strip():
            self._parts.append(data)
