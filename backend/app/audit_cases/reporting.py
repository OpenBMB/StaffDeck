from __future__ import annotations

import re
from collections import defaultdict
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from io import BytesIO

from docx import Document
from pydantic import BaseModel, Field
from sqlmodel import Session, select

from app.audit_cases.coverage import calculate_coverage, require_publishable
from app.audit_cases.elements import load_required_elements
from app.audit_cases.storage import read_case_blob, write_case_blob
from app.db.models import (
    AuditCase,
    AuditCaseDocument,
    AuditCaseDocumentVersion,
    AuditCaseMaterial,
    AuditEvidenceLedger,
    AuditReportSection,
    AuditReportVersion,
    ModelConfig,
    ProjectRuleBinding,
    RuleDefinition,
    RuleSetVersion,
)
from app.llm.client import LLMClient, LLMError

REPORT_SECTION_PROMPT = """
你是审核报告章节撰写器。只使用输入证据台账中的事实，不得补写、推测或引用未提供的资料。
每项事实必须使用 [EVIDENCE:<id>] 标注来源；只返回章节正文 Markdown。
""".strip()
EVIDENCE_CITATION = re.compile(r"\[EVIDENCE:([A-Za-z0-9_-]+)\]")
TEMPLATE_ANCHOR = re.compile(r"\{\{SECTION:([A-Za-z0-9_-]+)\}\}")


ClientFactory = Callable[[ModelConfig], LLMClient]


class ReportSectionSpec(BaseModel):
    id: str
    title: str
    sequence: int
    audit_element_ids: list[str]


class ReportGenerationSummary(BaseModel):
    status: str = "succeeded"
    generated_section_ids: list[str] = Field(default_factory=list)
    regenerated_section_ids: list[str] = Field(default_factory=list)


class AuditCitationError(RuntimeError):
    pass


class AuditReportBlocked(RuntimeError):
    pass


@dataclass(frozen=True)
class ReportRuleSnapshot:
    version_ids: list[str] = field(default_factory=list)
    definition_ids_by_section: dict[str, list[str]] = field(default_factory=dict)
    definitions: list[dict[str, object]] = field(default_factory=list)
    status: str = "not_configured"


class AuditReportService:
    def __init__(self, db: Session, client_factory: ClientFactory = LLMClient) -> None:
        self.db = db
        self.client_factory = client_factory

    def _document_scope(
        self,
        case: AuditCase,
        document_id: str | None,
        document_version_id: str | None,
    ) -> tuple[str | None, str | None]:
        if document_id is None:
            if document_version_id is not None:
                raise AuditReportBlocked("REPORT_DOCUMENT_SCOPE_REQUIRED")
            return None, None
        document = self.db.exec(
            select(AuditCaseDocument).where(
                AuditCaseDocument.id == document_id,
                AuditCaseDocument.tenant_id == case.tenant_id,
                AuditCaseDocument.audit_case_id == case.id,
            )
        ).first()
        if document is None:
            raise AuditReportBlocked("REPORT_DOCUMENT_NOT_FOUND")
        if document.status == "archived":
            raise AuditReportBlocked("REPORT_DOCUMENT_ARCHIVED")
        if not document.active_version_id:
            raise AuditReportBlocked("REPORT_DOCUMENT_VERSION_REQUIRED")
        selected_version_id = document_version_id or document.active_version_id
        version = self.db.exec(
            select(AuditCaseDocumentVersion).where(
                AuditCaseDocumentVersion.id == selected_version_id,
                AuditCaseDocumentVersion.tenant_id == case.tenant_id,
                AuditCaseDocumentVersion.audit_case_id == case.id,
                AuditCaseDocumentVersion.document_id == document.id,
            )
        ).first()
        if version is None:
            raise AuditReportBlocked("REPORT_DOCUMENT_VERSION_REQUIRED")
        if version.id != document.active_version_id:
            raise AuditReportBlocked("REPORT_DOCUMENT_VERSION_STALE")
        return document.id, version.id

    def rule_snapshot(
        self,
        case: AuditCase,
        source_document_id: str | None = None,
        source_document_version_id: str | None = None,
    ) -> ReportRuleSnapshot:
        source_document_id, _source_document_version_id = self._document_scope(
            case, source_document_id, source_document_version_id
        )
        statement = select(ProjectRuleBinding).where(
            ProjectRuleBinding.tenant_id == case.tenant_id,
            ProjectRuleBinding.audit_case_id == case.id,
            ProjectRuleBinding.status == "current",
        )
        if source_document_id is None:
            statement = statement.where(ProjectRuleBinding.document_id.is_(None))
        else:
            statement = statement.where(
                ProjectRuleBinding.document_id == source_document_id,
                ProjectRuleBinding.document_version_id == _source_document_version_id,
            )
        bindings = list(
            self.db.exec(
                statement.order_by(ProjectRuleBinding.priority, ProjectRuleBinding.bound_at)
            ).all()
        )
        section_ids = [spec.id for spec in _report_section_specs(case.management_systems_json)]
        section_rules = {section_id: [] for section_id in section_ids}
        if not bindings:
            return ReportRuleSnapshot(definition_ids_by_section=section_rules)

        version_ids: list[str] = []
        definitions: list[RuleDefinition] = []
        incomplete = False
        for binding in bindings:
            version = self.db.get(RuleSetVersion, binding.rule_set_version_id)
            if (
                version is None
                or version.tenant_id != case.tenant_id
                or version.status != "published"
            ):
                incomplete = True
                continue
            version_ids.append(version.id)
            definitions.extend(
                self.db.exec(
                    select(RuleDefinition)
                    .where(
                        RuleDefinition.tenant_id == case.tenant_id,
                        RuleDefinition.rule_set_version_id == version.id,
                        RuleDefinition.enabled,
                    )
                ).all()
            )

        applicable = [
            rule
            for rule in definitions
            if _rule_applies_to_report(rule)
        ]
        applicable.sort(key=lambda rule: (rule.rule_key, rule.sequence, rule.id))
        definition_payloads = [_rule_public_payload(rule) for rule in applicable]
        definition_ids = [rule.id for rule in applicable]
        section_rules = {section_id: list(definition_ids) for section_id in section_ids}
        return ReportRuleSnapshot(
            version_ids=sorted(set(version_ids)),
            definition_ids_by_section=section_rules,
            definitions=definition_payloads,
            status="incomplete" if incomplete else "complete",
        )

    def create_version(
        self,
        case: AuditCase,
        *,
        source_document_id: str | None = None,
        source_document_version_id: str | None = None,
    ) -> AuditReportVersion:
        latest = self.db.exec(
            select(AuditReportVersion)
            .where(
                AuditReportVersion.tenant_id == case.tenant_id,
                AuditReportVersion.audit_case_id == case.id,
            )
            .order_by(AuditReportVersion.version.desc())
        ).first()
        version = (latest.version + 1) if latest else 1
        source_document_id, source_document_version_id = self._document_scope(
            case, source_document_id, source_document_version_id
        )
        rule_snapshot = self.rule_snapshot(
            case, source_document_id, source_document_version_id
        )
        materials = list(
            self.db.exec(
                select(AuditCaseMaterial).where(
                    AuditCaseMaterial.tenant_id == case.tenant_id,
                    AuditCaseMaterial.audit_case_id == case.id,
                    AuditCaseMaterial.is_current,
                )
            ).all()
        )
        report = AuditReportVersion(
            tenant_id=case.tenant_id,
            audit_case_id=case.id,
            source_document_id=source_document_id,
            source_document_version_id=source_document_version_id,
            version=version,
            material_version_ids_json=[f"{row.id}:v{row.version}" for row in materials],
            knowledge_base_version_ids_json=list(case.knowledge_base_version_ids_json),
            rule_set_version_ids_json=list(rule_snapshot.version_ids),
            rule_traceability_status=rule_snapshot.status,
        )
        self.db.add(report)
        self.db.flush()
        for spec in _report_section_specs(case.management_systems_json):
            self.db.add(
                AuditReportSection(
                    tenant_id=case.tenant_id,
                    audit_case_id=case.id,
                    report_version_id=report.id,
                    section_id=spec.id,
                    title=spec.title,
                    sequence=spec.sequence,
                    audit_element_ids_json=list(spec.audit_element_ids),
                    rule_definition_ids_json=list(
                        rule_snapshot.definition_ids_by_section.get(spec.id, [])
                    ),
                )
            )
        self.db.commit()
        self.db.refresh(report)
        return report

    def generate_pending_sections(
        self, case: AuditCase, report: AuditReportVersion, model_config: ModelConfig
    ) -> ReportGenerationSummary:
        generated_section_ids: list[str] = []
        regenerated_section_ids: list[str] = []
        failed = False
        for section in self._pending_sections(report.id):
            generated_section_ids.append(section.section_id)
            if section.status == "failed":
                regenerated_section_ids.append(section.section_id)
            evidence = self._section_evidence(case.id, section.audit_element_ids_json)
            try:
                draft = self.client_factory(model_config).generate_text(
                    REPORT_SECTION_PROMPT,
                    {
                        "section_id": section.section_id,
                        "title": section.title,
                        "evidence": [ledger_public_payload(item) for item in evidence],
                        "rules": self._section_rules(report, section),
                        "citation_rule": "每项事实使用 [EVIDENCE:<id>] 引用",
                    },
                )
                citation_ids = validate_section_citations(draft, evidence)
                section.draft_markdown = draft
                section.citation_ids_json = citation_ids
                section.status = "succeeded"
                section.error_code = None
            except (LLMError, AuditCitationError) as exc:
                failed = True
                section.status = "failed"
                section.retry_count += 1
                section.error_code = (
                    "REPORT_CITATION_INVALID"
                    if isinstance(exc, AuditCitationError)
                    else "REPORT_MODEL_FAILED"
                )
            self.db.add(section)
            self.db.commit()
        return ReportGenerationSummary(
            status="failed" if failed else "succeeded",
            generated_section_ids=generated_section_ids,
            regenerated_section_ids=regenerated_section_ids,
        )

    def assemble_draft(self, report: AuditReportVersion) -> str:
        sections = self._sections(report.id)
        if any(row.status != "succeeded" for row in sections):
            raise AuditReportBlocked("REPORT_SECTIONS_INCOMPLETE")
        return "\n\n".join(
            f"## {row.title}\n\n{row.draft_markdown}"
            for row in sorted(sections, key=lambda item: item.sequence)
        )

    def publish(
        self, case: AuditCase, report: AuditReportVersion, confirmed_by: str | None
    ) -> AuditReportVersion:
        if confirmed_by and report.rule_traceability_status != "complete":
            raise AuditReportBlocked("RULE_BINDING_REQUIRED_FOR_PUBLISH")
        snapshot = calculate_coverage(self.db, case)
        require_publishable(snapshot)
        markdown = self.assemble_draft(report)
        filename = "审核报告.docx" if confirmed_by else "审核报告-待确认草稿.docx"
        template = self.db.exec(
            select(AuditCaseMaterial).where(
                AuditCaseMaterial.tenant_id == case.tenant_id,
                AuditCaseMaterial.audit_case_id == case.id,
                AuditCaseMaterial.material_type == "report_template",
                AuditCaseMaterial.is_current,
            )
        ).first()
        if template is None:
            data = render_docx(markdown, title=filename.removesuffix(".docx"))
        else:
            sections = {
                row.section_id: row.draft_markdown for row in self._sections(report.id)
            }
            data = render_docx_from_template(read_case_blob(template.storage_key), sections)
        report.final_storage_key = write_case_blob(
            tenant_id=case.tenant_id,
            audit_case_id=case.id,
            material_id=report.id,
            name=filename,
            data=data,
        )
        report.coverage_snapshot_json = snapshot.model_dump(mode="json")
        report.status = "published" if confirmed_by else "review"
        report.lead_auditor_confirmed_by = confirmed_by
        self.db.add(report)
        self.db.commit()
        self.db.refresh(report)
        return report

    def _pending_sections(self, report_id: str) -> list[AuditReportSection]:
        return list(
            self.db.exec(
                select(AuditReportSection)
                .where(
                    AuditReportSection.report_version_id == report_id,
                    AuditReportSection.status.in_(["pending", "failed"]),
                )
                .order_by(AuditReportSection.sequence)
            ).all()
        )

    def _sections(self, report_id: str) -> list[AuditReportSection]:
        return list(
            self.db.exec(
                select(AuditReportSection)
                .where(AuditReportSection.report_version_id == report_id)
                .order_by(AuditReportSection.sequence)
            ).all()
        )

    def _section_rules(
        self, report: AuditReportVersion, section: AuditReportSection
    ) -> list[dict[str, object]]:
        allowed_ids = set(section.rule_definition_ids_json or [])
        if not allowed_ids:
            return []
        rules = list(
            self.db.exec(
                select(RuleDefinition).where(
                    RuleDefinition.tenant_id == report.tenant_id,
                    RuleDefinition.id.in_(allowed_ids),
                )
            ).all()
        )
        rules.sort(key=lambda rule: (rule.rule_key, rule.sequence, rule.id))
        return [_rule_public_payload(rule) for rule in rules]

    def _section_evidence(
        self, case_id: str, audit_element_ids: list[str]
    ) -> list[AuditEvidenceLedger]:
        if not audit_element_ids:
            return []
        return list(
            self.db.exec(
                select(AuditEvidenceLedger)
                .where(
                    AuditEvidenceLedger.audit_case_id == case_id,
                    AuditEvidenceLedger.audit_element_id.in_(audit_element_ids),
                )
                .order_by(AuditEvidenceLedger.audit_element_id, AuditEvidenceLedger.created_at)
            ).all()
        )


def _report_section_specs(management_systems: list[str]) -> list[ReportSectionSpec]:
    grouped: dict[str, list[str]] = defaultdict(list)
    titles: dict[str, str] = {}
    for element in load_required_elements(management_systems):
        if element.report_section_id not in titles:
            titles[element.report_section_id] = element.title
        grouped[element.report_section_id].append(element.id)
    return [
        ReportSectionSpec(
            id=section_id,
            title=titles[section_id],
            sequence=sequence,
            audit_element_ids=sorted(element_ids),
        )
        for sequence, (section_id, element_ids) in enumerate(sorted(grouped.items()), start=1)
    ]


def _rule_applies_to_report(rule: RuleDefinition) -> bool:
    workflow_nodes = set(rule.workflow_nodes_json or [])
    document_types = set(rule.document_types_json or [])
    workflow_match = not workflow_nodes or "generate_report_sections" in workflow_nodes
    document_match = not document_types or "audit_report" in document_types
    return workflow_match and document_match


def _rule_public_payload(rule: RuleDefinition) -> dict[str, object]:
    return {
        "id": rule.id,
        "rule_key": rule.rule_key,
        "name": rule.name,
        "execution_level": rule.execution_level,
        "execution_method": rule.execution_method,
        "workflow_nodes": list(rule.workflow_nodes_json or []),
        "information_domains": list(rule.information_domains_json or []),
        "document_types": list(rule.document_types_json or []),
        "source_refs": [dict(item) for item in (rule.source_refs_json or [])],
    }


def ledger_public_payload(row: AuditEvidenceLedger) -> dict[str, object]:
    return {
        "id": row.id,
        "audit_element_id": row.audit_element_id,
        "source_kind": row.source_kind,
        "source_ref": row.source_ref,
        "source_version_id": row.source_version_id,
        "evidence_type": row.evidence_type,
        "evidence_text": row.evidence_text,
        "confidence": row.confidence,
    }


def validate_section_citations(
    markdown: str, evidence: list[AuditEvidenceLedger]
) -> list[str]:
    allowed = {item.id for item in evidence}
    cited = EVIDENCE_CITATION.findall(markdown)
    if not cited or any(item not in allowed for item in cited):
        raise AuditCitationError("REPORT_CITATION_INVALID")
    return list(dict.fromkeys(cited))


def render_docx(markdown: str, *, title: str) -> bytes:
    document = Document()
    document.add_heading(title, level=0)
    for block in markdown.split("\n\n"):
        if block.startswith("## "):
            document.add_heading(block.removeprefix("## "), level=1)
        elif block.strip():
            document.add_paragraph(block)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def render_docx_from_template(
    template_bytes: bytes, sections: dict[str, str]
) -> bytes:
    document = Document(BytesIO(template_bytes))
    paragraphs = list(_iter_paragraphs(document))
    document_text = "\n".join(paragraph.text for paragraph in paragraphs)
    anchors = TEMPLATE_ANCHOR.findall(document_text)
    expected = set(sections)
    if len(anchors) != len(expected) or set(anchors) != expected or len(anchors) != len(set(anchors)):
        raise AuditReportBlocked("REPORT_TEMPLATE_ANCHOR_MISSING")
    for paragraph in paragraphs:
        for section_id, markdown in sections.items():
            anchor = f"{{{{SECTION:{section_id}}}}}"
            if anchor in paragraph.text:
                paragraph.text = paragraph.text.replace(anchor, markdown)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _iter_paragraphs(document: Document) -> Iterable[object]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
