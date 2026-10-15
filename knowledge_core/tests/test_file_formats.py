import base64
import time
from io import BytesIO

from fastapi.testclient import TestClient
from docx import Document

from knowledge_core.api import create_app
from support import full_runtime as Runtime


def _wait_job(client, job_id):
    for _ in range(200):
        job = client.get(f"/api/km/v1/jobs/{job_id}").json()
        if job["status"] in {"succeeded", "failed", "cancelled"}:
            return job
        time.sleep(0.01)
    raise AssertionError("ingest did not finish")


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Expense policy", level=1)
    document.add_paragraph("DOCX receipts must be retained for 30 days.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes() -> bytes:
    # Small self-contained PDF fixture so the acceptance test does not depend
    # on a document-generation package outside the KM wheel.
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length 68 >>\nstream\nBT /F1 12 Tf 72 720 Td (PDF receipts require manager approval.) Tj ET\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def test_html_pdf_docx_uploads_preserve_real_text_and_type(tmp_path):
    runtime = Runtime(f"sqlite:///{tmp_path / 'formats.sqlite'}")
    fixtures = {
        "policy.html": ("html", b"<html><body><h1>HTML policy</h1><p>HTML receipts need approval.</p></body></html>", "HTML receipts"),
        "policy.pdf": ("pdf", _pdf_bytes(), "PDF receipts"),
        "policy.docx": ("docx", _docx_bytes(), "DOCX receipts"),
    }
    with TestClient(create_app(runtime)) as client:
        for filename, (file_type, content, query) in fixtures.items():
            base = client.post("/api/km/v1/knowledge-bases", json={"name": filename}).json()
            response = client.post(
                f"/api/km/v1/knowledge-bases/{base['id']}/documents",
                json={"filename": filename, "content_base64": base64.b64encode(content).decode()},
            )
            assert response.status_code == 202, response.text
            job = _wait_job(client, response.json()["id"])
            assert job["status"] == "succeeded", job
            document = client.get(f"/api/km/v1/documents/{job['document_id']}").json()
            assert document["file_type"] == file_type
            assert query in document["metadata"]["raw_text"]
            result = client.post("/api/km/v1/query", json={"query": query}).json()
            assert result["chunks"], (filename, result)


def test_minimal_html_pdf_docx_ingest_uses_parser_type_and_needs_no_catalog(tmp_path):
    from knowledge_core.runtime import Runtime as ModularRuntime
    runtime = ModularRuntime(f"sqlite:///{tmp_path / 'minimal-formats.sqlite'}")
    fixtures = {
        'policy.HTML': ('html', b'<h1>HTML policy</h1><p>HTML receipts need approval.</p>', 'HTML receipts'),
        'policy.pdf': ('pdf', _pdf_bytes(), 'PDF receipts'),
        'policy.docx': ('docx', _docx_bytes(), 'DOCX receipts'),
    }
    with TestClient(create_app(runtime)) as client:
        for filename, (kind, content, query) in fixtures.items():
            upload = client.post('/api/km/v1/documents', json={'filename': filename, 'content_base64': base64.b64encode(content).decode()})
            assert upload.status_code == 202
            job = _wait_job(client, upload.json()['id']); assert job['status'] == 'succeeded', job
            document = client.get('/api/km/v1/documents/' + job['document_id']).json()
            assert document['file_type'] == kind and query in document['content']
            assert client.post('/api/km/v1/query', json={'query': query}).json()['chunks']
